#!/usr/bin/env python3
"""Import the two supplied Stapfer XLSX exports without modifying legacy sources.

Requires openpyxl. No macros, formulas, links, or workbook commands are executed.
Original workbooks remain external inputs; generated records carry file/row hashes.
"""
import argparse
import collections
import hashlib
import json
import math
import pathlib
import re
import zipfile

HEADERS = ["Schulort", "Schule Konfession", "Schultyp", "Frage", "Antwort",
           "Signatur", "TransID", "Distrikt", "Kanton 2015", "Kanton 1799",
           "Kanton 1750", "Zitierempfehlung"]
DATASET = "stapfer-1799-answers"
URL = "https://opendata.swiss/de/dataset/stapfer-enquete-schule-in-der-schweiz-17991"
LICENCE = "https://opendata.swiss/de/terms-of-use#terms_open"


def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


def digest(value):
    return hashlib.sha256(value).hexdigest()


def assess(values):
    """Return a stable rejection reason. Preserve empty-answer rows as missingness."""
    if len(values) != 12:
        return "column_count"
    if any(not (v is None or isinstance(v, (str, int, float, bool))) for v in values):
        return "unsupported_cell_type"
    if any(isinstance(v, str) and v.startswith("=") for v in values):
        return "formula_cell"
    if not re.fullmatch(r"[0-9]{1,12}", str(values[6] or "")):
        return "invalid_transcription_id"
    if not isinstance(values[3], str) or not values[3].strip():
        return "missing_question"
    if values[4] is None or (isinstance(values[4], str) and not values[4].strip()):
        return "empty_answer"
    if isinstance(values[4], bool) or not isinstance(values[4], (str, int, float)):
        return "nontext_answer"
    if isinstance(values[4], float) and not math.isfinite(values[4]):
        return "nonfinite_answer"
    if len(answer_text(values[4]).encode("utf-8")) > 100_000:
        return "answer_too_large"
    return None


def answer_text(value):
    return value if isinstance(value, str) else canonical(value)


def read_workbook(path):
    import openpyxl
    if path.stat().st_size > 16 * 1024 * 1024:
        raise ValueError("workbook_byte_limit")
    with zipfile.ZipFile(path) as package:
        if len(package.infolist()) > 500 or sum(i.file_size for i in package.infolist()) > 128 * 1024 * 1024:
            raise ValueError("workbook_expansion_limit")
        if any("vbaProject" in i.filename or "/externalLinks/" in i.filename for i in package.infolist()):
            raise ValueError("active_workbook_content")
        for entry in package.infolist():
            if entry.filename.endswith((".xml", ".rels")):
                markup = package.read(entry).upper()
                if b"<!DOCTYPE" in markup or b"<!ENTITY" in markup:
                    raise ValueError("workbook_xml_entities")
    file_hash = digest(path.read_bytes())
    book = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if book.sheetnames != ["Worksheet"]:
            raise ValueError("unexpected_worksheets")
        sheet = book.active
        if sheet.max_row > 25_001 or sheet.max_column > 32:
            raise ValueError("workbook_dimension_limit")
        iterator = sheet.iter_rows(values_only=True)
        headers = list(next(iterator))
        if headers[:12] != HEADERS or any(v is not None for v in headers[12:]):
            raise ValueError("unexpected_headers")
        rows = []
        for number, full in enumerate(iterator, 2):
            values = list(full[:12])
            reason = "unmapped_nonempty_column" if any(v is not None for v in full[12:]) else assess(values)
            # Reject unsupported cells before canonical serialization, retaining a typed representation.
            raw = [v if v is None or isinstance(v, (str, int, float, bool)) else
                   {"unsupported_type": type(v).__name__, "value": str(v)} for v in values]
            row = {"row": number, "locator": f"{path.name}#Worksheet!A{number}:L{number}",
                   "row_sha256": digest(canonical(raw).encode()), "source_sha256": file_hash,
                   "values": values, "reason": reason}
            rows.append(row)
        return {"file": path.name, "sha256": file_hash, "bytes": path.stat().st_size,
                "sheet": "Worksheet", "rows": rows}
    finally:
        book.close()


def rejection(row):
    return {"locator": row["locator"], "row": row["row"], "row_sha256": row["row_sha256"],
            "source_sha256": row["source_sha256"], "reason": row["reason"],
            "transcription_id": str(row["values"][6]) if row["values"][6] is not None else None,
            "question": row["values"][3], "content_retained_in_original": True}


def group_sources(workbook):
    groups = collections.defaultdict(list)
    for row in workbook["rows"]:
        if row["reason"] in (None, "empty_answer"):
            groups[str(row["values"][6])].append(row)
    sources = []
    for trans_id, rows in sorted(groups.items(), key=lambda item: int(item[0])):
        # A TransID is not assumed to identify one school if exported metadata conflicts.
        stable = [0, 1, 2, 5, 7, 8, 9, 10, 11]
        if any(len({canonical(r["values"][i]) for r in rows}) > 1 for i in stable):
            for row in rows:
                row["reason"] = "conflicting_transcription_metadata"
            continue
        metadata = dict(zip(HEADERS, rows[0]["values"]))
        metadata.pop("Antwort"); metadata.pop("Frage"); metadata["TransID"] = trans_id
        text, provenance, missing = "", [], []
        for row in rows:
            item = {k: row[k] for k in ("row", "locator", "row_sha256")}
            item["question"] = row["values"][3]
            if row["reason"]:
                item["reason"] = row["reason"]
                missing.append(item)
            else:
                text += ("\n\n" if text else "") + f"[Frage: {row['values'][3]}]\n"
                start = len(text)
                answer = answer_text(row["values"][4])
                text += answer
                item.update({"answer_start_char": start, "answer_end_char": len(text),
                             "answer_sha256": digest(answer.encode()),
                             "answer_cell_type": type(row["values"][4]).__name__})
                if not isinstance(row["values"][4], str):
                    item["original_numeric_value"] = row["values"][4]
                provenance.append(item)
        row_manifest = [{"row": r["row"], "row_sha256": r["row_sha256"]} for r in rows]
        source = {"id": f"stapfer-answer-{trans_id}", "dataset_id": DATASET,
                  "title": f"{metadata['Schulort']} — Antworten (TransID {trans_id})", "text": text,
                  "normalized_text": text.replace("\r\n", "\n").replace("\r", "\n"),
                  "url": URL, "page": None, "licence": LICENCE, "metadata": metadata,
                  "sha256": digest(canonical(row_manifest).encode()),
                  "record_hash_scope": "canonical-import-row-manifest-json",
                  "text_sha256": digest(text.encode()), "source_sha256": workbook["sha256"],
                  "record_locator": {"kind": "xlsx-transcription-group", "file": workbook["file"],
                                     "sheet": "Worksheet", "transcription_id": trans_id},
                  "provenance": {"accepted_rows": provenance, "missing_rows": missing,
                                 "expected_row_count": len(rows), "accepted_row_count": len(provenance),
                                 "missing_row_count": len(missing), "row_manifest": row_manifest},
                  "transformation": {"kind": "question-answer-row-assembly", "version": "1",
                                     "text_answer_bytes_modified": False, "numeric_answer_rendering": "canonical JSON scalar; original numeric type/value retained",
                                     "separator": "question labels and two newlines",
                                     "offset_unit": "Unicode code points", "original_files_bundled": False},
                  "flags": {"complete_transcription_claim": False, "page_reference_missing": True,
                            "missing_answers": len(missing), "handwriting_status": "not-assessed",
                            "historical_accuracy_human_review": "NOT_RUN"}}
        sources.append(source)
    return sources


def run(directory, legacy, output):
    workbooks = [read_workbook(directory / name) for name in ("export_antworten.xlsx", "export_antworten(1).xlsx")]
    existing = [json.loads(line) for line in legacy.read_text().splitlines()]
    by_id = {r["id"]: r for r in existing}
    legacy_rows = []
    for row in workbooks[0]["rows"]:
        if row["reason"]:
            raise ValueError("legacy_workbook_contains_rejected_row")
        old = by_id.get(f"stapfer-{row['values'][6]}")
        if not old or old["text"] != row["values"][4] or old["source_sha256"] != workbooks[0]["sha256"]:
            raise ValueError("legacy_source_drift")
        legacy_rows.append({"id": old["id"], "legacy_record_sha256": old["sha256"],
                            "text_sha256": digest(old["text"].encode()),
                            **{k: row[k] for k in ("locator", "row", "row_sha256")}})
    if len(existing) != len(legacy_rows):
        raise ValueError("legacy_count_drift")
    sources = group_sources(workbooks[1])
    rejected = [rejection(row) for workbook in workbooks for row in workbook["rows"] if row["reason"]]
    manifest = {"schema_version": "1.0", "row_hash_scope": "canonical JSON array of first 12 raw cell values",
                "canonical_json": "UTF-8, ensure_ascii=false, sorted object keys, compact separators, no NaN",
                "files": [{k: w[k] for k in ("file", "sha256", "bytes", "sheet")} |
                          {"data_rows": len(w["rows"]), "accepted_rows": sum(r["reason"] is None for r in w["rows"]),
                           "rejected_rows": sum(r["reason"] is not None for r in w["rows"])} for w in workbooks],
                "legacy_source_file_sha256": digest(legacy.read_bytes()), "legacy_sources_unchanged": True,
                "legacy_row_audit": legacy_rows, "second_collection_id": DATASET,
                "second_collection_records": len(sources),
                "overlapping_transcription_ids": len(set(str(r["values"][6]) for r in workbooks[0]["rows"]) &
                                                      set(s["metadata"]["TransID"] for s in sources)),
                "rejection_reasons": dict(collections.Counter(r["reason"] for r in rejected)),
                "original_files_committed": False, "provider_calls": 0,
                "limitations": ["Workbook rows are published transcription exports, not original-page OCR.",
                                "Question labels are source export labels; numeric relational question IDs are not supplied.",
                                "Blank answers indicate missing exported content, not negative historical evidence.",
                                "Collections overlap and must not be counted as independent historical observations."]}
    output.mkdir(parents=True, exist_ok=True)
    for name, values in [("stapfer-answer-sources.jsonl", sources), ("stapfer-rejected-rows.jsonl", rejected)]:
        (output / name).write_text("".join(canonical(v) + "\n" for v in values))
    manifest["outputs"] = [{"path": name, "sha256": digest((output / name).read_bytes())}
                           for name in ("stapfer-answer-sources.jsonl", "stapfer-rejected-rows.jsonl")]
    (output / "stapfer-import-manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    return manifest


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=pathlib.Path, required=True)
    parser.add_argument("--legacy", type=pathlib.Path, default=pathlib.Path("data/stapfer-sources.jsonl"))
    parser.add_argument("--output", type=pathlib.Path, default=pathlib.Path("data"))
    args = parser.parse_args()
    receipt = run(args.input_dir, args.legacy, args.output)
    print(json.dumps({"records": receipt["second_collection_records"], "rejections": receipt["rejection_reasons"],
                      "legacy_sources_unchanged": receipt["legacy_sources_unchanged"]}))
